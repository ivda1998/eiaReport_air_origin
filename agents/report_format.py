"""
ReportFormatAgent: 보고서 저장 및 DOCX 스타일링

역할:
- TXT 보고서 저장
- DOCX 보고서 생성 (마크다운 → DOCX 변환)
- DOCX 스타일링 (표지, 머리글/바닥글, 표 스타일)
- HWP 변환 안내
"""

import os
import time

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    Document = None


# ==========================================
# DOCX 스타일링 헬퍼 함수
# ==========================================

def _parse_pt(value):
    """'16pt' 같은 문자열에서 숫자 추출"""
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str):
        return int(value.replace('pt', '').replace('mm', '').strip())
    return 10


def _parse_color(hex_color):
    """'#2E5090' -> RGBColor"""
    if not hex_color or not isinstance(hex_color, str):
        return RGBColor(0x21, 0x21, 0x21)
    hex_color = hex_color.lstrip('#')
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def _setup_page(doc, styles_config):
    """페이지 설정 (여백)"""
    page_config = styles_config.get('document', {}).get('page', {})
    if not page_config:
        page_config = styles_config.get('page', {})
    margins = page_config.get('margins', {})

    section = doc.sections[0]
    if margins:
        section.top_margin = Mm(int(str(margins.get('top', '25')).replace('mm', '')))
        section.bottom_margin = Mm(int(str(margins.get('bottom', '25')).replace('mm', '')))
        section.left_margin = Mm(int(str(margins.get('left', '30')).replace('mm', '')))
        section.right_margin = Mm(int(str(margins.get('right', '30')).replace('mm', '')))


def _setup_heading_styles(doc, styles_config):
    """H1~H3 제목 스타일 설정"""
    heading_config = styles_config.get('heading_styles', {})
    if not heading_config:
        heading_config = styles_config.get('headings', {})
    if not heading_config:
        return

    for level, key in [(1, 'h1'), (2, 'h2'), (3, 'h3')]:
        h_conf = heading_config.get(key, {})
        if not h_conf:
            continue
        try:
            style = doc.styles[f'Heading {level}']
            font = style.font
            font.name = h_conf.get('font_family', h_conf.get('font_name', '맑은 고딕'))
            font.size = Pt(_parse_pt(h_conf.get('font_size', 16 - (level-1)*2)))
            font.bold = h_conf.get('font_weight') == 'bold'
            font.color.rgb = _parse_color(h_conf.get('color', '#212121'))

            pf = style.paragraph_format
            pf.space_before = Pt(_parse_pt(h_conf.get('spacing_before', 12)))
            pf.space_after = Pt(_parse_pt(h_conf.get('spacing_after', 6)))
        except Exception:
            pass


def _setup_body_style(doc, styles_config):
    """본문 스타일 설정"""
    body_config = styles_config.get('body_styles', {}).get('normal', {})
    if not body_config:
        body_config = styles_config.get('body', {})
    if not body_config:
        return

    try:
        style = doc.styles['Normal']
        font = style.font
        font.name = body_config.get('font_family', '맑은 고딕')
        font.size = Pt(_parse_pt(body_config.get('font_size', 10)))
        font.color.rgb = _parse_color(body_config.get('color', '#212121'))

        pf = style.paragraph_format
        spacing = body_config.get('line_spacing', 1.6)
        if isinstance(spacing, (int, float)):
            pf.line_spacing = spacing
        pf.space_after = Pt(_parse_pt(body_config.get('spacing_after', 6)))
    except Exception:
        pass


def _add_cover_page(doc, styles_config, project_config):
    """표지 페이지 생성"""
    cover = styles_config.get('cover', styles_config.get('cover_fonts', {}))
    if not cover:
        return

    # 빈 줄 추가로 표지 상단 여백
    for _ in range(6):
        doc.add_paragraph()

    # 제목
    title_conf = cover.get('title', {})
    title_text = project_config.get('name', '환경영향평가서')
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title_text)
    title_run.font.size = Pt(_parse_pt(title_conf.get('font_size', 24)))
    title_run.font.bold = title_conf.get('font_weight') == 'bold'
    title_run.font.color.rgb = _parse_color(title_conf.get('color', '#2E5090'))
    title_run.font.name = title_conf.get('font_family', '맑은 고딕')

    # 부제목
    subtitle_conf = cover.get('subtitle', {})
    subtitle_para = doc.add_paragraph()
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_para.add_run("환경영향평가서 - 대기질 부문")
    subtitle_run.font.size = Pt(_parse_pt(subtitle_conf.get('font_size', 18)))
    subtitle_run.font.color.rgb = _parse_color(subtitle_conf.get('color', '#4A7BA7'))
    subtitle_run.font.name = subtitle_conf.get('font_family', '맑은 고딕')

    # 빈 줄
    for _ in range(4):
        doc.add_paragraph()

    # 프로젝트 정보
    info_conf = cover.get('info', {})
    info_items = [
        f"사업명: {project_config.get('name', '')}",
        f"작성일: {project_config.get('date', time.strftime('%Y-%m-%d'))}",
        f"작성 도구: AERMOD + AI 자동 작성 시스템",
    ]
    for item in info_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(item)
        r.font.size = Pt(_parse_pt(info_conf.get('font_size', 12)))
        r.font.name = info_conf.get('font_family', '맑은 고딕')
        r.font.color.rgb = _parse_color(info_conf.get('color', '#212121'))

    # 페이지 나누기
    doc.add_page_break()


def _add_header_footer(doc, styles_config, project_name):
    """머리글/바닥글 설정"""
    try:
        header_conf = styles_config.get('header', {})
        footer_conf = styles_config.get('footer', {})

        section = doc.sections[0]

        # 머리글
        if header_conf.get('enabled', False):
            header = section.header
            header.is_linked_to_previous = False
            p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            p.text = ""
            run_left = p.add_run(f"{project_name}")
            run_left.font.size = Pt(_parse_pt(header_conf.get('font_size', 9)))
            run_left.font.color.rgb = _parse_color(header_conf.get('color', '#757575'))
            run_left.font.name = header_conf.get('font_family', '맑은 고딕')
            run_sep = p.add_run("    |    ")
            run_sep.font.size = Pt(9)
            run_sep.font.color.rgb = _parse_color('#BDBDBD')
            run_right = p.add_run("환경영향평가서")
            run_right.font.size = Pt(_parse_pt(header_conf.get('font_size', 9)))
            run_right.font.color.rgb = _parse_color(header_conf.get('color', '#757575'))
            run_right.font.name = header_conf.get('font_family', '맑은 고딕')

        # 바닥글
        if footer_conf.get('enabled', False):
            footer = section.footer
            footer.is_linked_to_previous = False
            p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"{project_name} | {time.strftime('%Y년 %m월 %d일')}")
            run.font.size = Pt(_parse_pt(footer_conf.get('font_size', 9)))
            run.font.color.rgb = _parse_color(footer_conf.get('color', '#757575'))
            run.font.name = footer_conf.get('font_family', '맑은 고딕')
    except Exception as e:
        print(f"  ⚠️ 머리글/바닥글 설정 실패: {e}")


def _style_table(table, styles_config):
    """표에 스타일 적용 (헤더 색상, 테두리 등)"""
    table_conf = styles_config.get('table_styles', {}).get('default', {})
    if not table_conf:
        table_conf = styles_config.get('tables', {}).get('default', {})
    if not table_conf:
        return

    header_conf = table_conf.get('header', {})
    bg_color = header_conf.get('background_color', '#2E5090').lstrip('#')
    text_color = header_conf.get('text_color', '#FFFFFF')

    # 첫 행(헤더)에 배경색 적용
    if table.rows:
        for cell in table.rows[0].cells:
            try:
                tc_pr = cell._tc.get_or_add_tcPr()
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), bg_color)
                shading.set(qn('w:val'), 'clear')
                tc_pr.append(shading)

                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = _parse_color(text_color)
                        run.font.size = Pt(_parse_pt(table_conf.get('font_size', 9)))
                        run.font.name = table_conf.get('font_family', '맑은 고딕')
            except Exception:
                pass

    # 데이터 행 스타일
    body_conf = table_conf.get('body', {})
    alt_color = body_conf.get('alternate_color', '#F5F5F5').lstrip('#')
    use_alt = body_conf.get('alternate_rows', False)

    for row_idx, row in enumerate(table.rows[1:], 1):
        for cell in row.cells:
            try:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(_parse_pt(table_conf.get('font_size', 9)))
                        run.font.name = table_conf.get('font_family', '맑은 고딕')

                if use_alt and row_idx % 2 == 0:
                    tc_pr = cell._tc.get_or_add_tcPr()
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), alt_color)
                    shading.set(qn('w:val'), 'clear')
                    tc_pr.append(shading)
            except Exception:
                pass


def _apply_docx_styles(doc, styles_config):
    """DOCX 문서에 전체 스타일 적용"""
    _setup_page(doc, styles_config)
    _setup_heading_styles(doc, styles_config)
    _setup_body_style(doc, styles_config)


# ==========================================
# 보고서 저장 함수
# ==========================================

def save_report(content, config, templates=None, format='txt', isopleth_images=None):
    """보고서를 파일로 저장

    Args:
        content: 보고서 텍스트
        config: ProjectConfig 인스턴스
        templates: TemplateManager 인스턴스 (DOCX 스타일링용)
        format: 'txt', 'docx', 'hwp'
        isopleth_images: 등농도곡선 이미지 경로 목록
    Returns:
        str or None: 저장된 파일 경로
    """
    if not os.path.exists(config.OUTPUT_DIR):
        os.makedirs(config.OUTPUT_DIR)

    if not content or not content.strip():
        print("⚠️ 저장할 내용이 비어있습니다.")
        return None

    if format == 'txt':
        report_file = os.path.join(config.OUTPUT_DIR, "대기질_환경영향평가서.txt")
        try:
            with open(report_file, "w", encoding="utf-8") as f:
                f.write(content)
                if isopleth_images:
                    f.write("\n\n[등농도곡선 이미지]\n")
                    for img_path in isopleth_images:
                        f.write(f"- {img_path}\n")
            print(f"\n💾 보고서 저장 완료: {report_file}")
            print(f"   파일 크기: {len(content):,}자")
            return report_file
        except Exception as e:
            print(f"❌ TXT 파일 저장 실패: {e}")
            return None

    elif format == 'docx':
        if Document is None:
            print("\n⚠️ python-docx가 설치되지 않았습니다. TXT 파일로 저장합니다.")
            return save_report(content, config, templates, 'txt', isopleth_images)

        timestamp = time.strftime("%Y%m%d_%H%M%S")
        report_file = os.path.join(config.OUTPUT_DIR, f"대기질_환경영향평가서_{timestamp}.docx")

        base_report_file = os.path.join(config.OUTPUT_DIR, "대기질_환경영향평가서.docx")
        if os.path.exists(base_report_file):
            print(f"\n⚠️ 기존 파일이 있습니다. 새 파일명으로 저장합니다: {os.path.basename(report_file)}")

        try:
            doc = Document()

            docx_config = config.config['report']['docx']

            # 템플릿 스타일 적용
            if templates is not None:
                docx_styles = templates.get_docx_styles()
                _apply_docx_styles(doc, docx_styles)
                _add_cover_page(doc, docx_styles, config.config['project'])
                _add_header_footer(doc, docx_styles, config.config['project']['name'])
            else:
                title = doc.add_heading(docx_config['title'], level=1)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            lines = content.split('\n')
            i = 0

            while i < len(lines):
                line = lines[i].strip()

                if not line:
                    i += 1
                    continue

                if line.startswith('|'):
                    table_lines = []
                    while i < len(lines) and lines[i].strip().startswith('|'):
                        table_lines.append(lines[i].strip())
                        i += 1

                    if len(table_lines) > 1:
                        try:
                            header_cells = [cell.strip() for cell in table_lines[0].split('|')[1:-1]]
                            num_cols = len(header_cells)

                            data_lines = [table_lines[0]]
                            for tl in table_lines[1:]:
                                if not tl.replace('|', '').replace('-', '').replace(' ', '').replace(':', ''):
                                    continue
                                data_lines.append(tl)

                            if len(data_lines) > 0:
                                table = doc.add_table(rows=len(data_lines), cols=num_cols)
                                table.style = 'Table Grid'

                                for row_idx, data_line in enumerate(data_lines):
                                    cells = [cell.strip() for cell in data_line.split('|')[1:-1]]
                                    for col_idx, cell_text in enumerate(cells):
                                        if col_idx < num_cols:
                                            cell = table.rows[row_idx].cells[col_idx]
                                            cell.text = cell_text

                                # 템플릿 스타일 적용
                                if templates is not None:
                                    _style_table(table, templates.get_docx_styles())
                        except Exception as e:
                            p = doc.add_paragraph('\n'.join(table_lines))
                            p.paragraph_format.line_spacing = 1.0

                    continue

                if line.startswith('#'):
                    level = min(line.count('#', 0, 4), 3)
                    text = line.lstrip('#').strip()
                    if text:
                        doc.add_heading(text, level=level)
                    i += 1
                    continue

                if line[0:3].replace('.', '').replace(')', '').replace('(', '').replace('가', '').replace('나', '').strip():
                    first_chars = line.split()[0] if line.split() else ""
                    if any(c in first_chars for c in ['.', ')', '(']):
                        doc.add_heading(line, level=2)
                        i += 1
                        continue

                para_lines = [line]
                i += 1
                while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith('#') and not lines[i].strip().startswith('|'):
                    para_lines.append(lines[i].strip())
                    i += 1

                para_text = ' '.join(para_lines)
                if para_text:
                    p = doc.add_paragraph(para_text)
                    p.paragraph_format.line_spacing = docx_config['line_spacing']

                    for run in p.runs:
                        run.font.name = docx_config['font_name']
                        run.font.size = Pt(docx_config['font_size'])

            if isopleth_images:
                doc.add_page_break()
                doc.add_heading('등농도곡선', level=1)
                for img_path in isopleth_images:
                    if os.path.exists(img_path):
                        try:
                            img_filename = os.path.basename(img_path)
                            doc.add_paragraph(f"[그림] {img_filename.replace('.png', '')}")
                            doc.add_picture(img_path, width=Inches(docx_config['image_width']))
                            last_paragraph = doc.paragraphs[-1]
                            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            doc.add_paragraph()
                        except Exception as e:
                            print(f"⚠️ 이미지 삽입 실패 ({img_filename}): {e}")

            max_retries = 3
            for attempt in range(max_retries):
                try:
                    doc.save(report_file)
                    print(f"\n💾 보고서 저장 완료: {report_file}")
                    print(f"   파일명: {os.path.basename(report_file)}")
                    print("   ℹ️ DOCX 파일은 한글(HWP)에서 열어서 HWP 형식으로 저장할 수 있습니다.")
                    return report_file
                except PermissionError as pe:
                    if attempt < max_retries - 1:
                        print(f"\n⚠️ 파일 저장 실패 (시도 {attempt+1}/{max_retries})")
                        print("   파일이 다른 프로그램에서 열려있을 수 있습니다.")
                        timestamp = time.strftime("%Y%m%d_%H%M%S") + f"_{attempt+1}"
                        report_file = os.path.join(config.OUTPUT_DIR, f"대기질_환경영향평가서_{timestamp}.docx")
                        print(f"   새 파일명으로 재시도: {os.path.basename(report_file)}")
                        time.sleep(1)
                    else:
                        raise pe

        except PermissionError as pe:
            print(f"\n❌ DOCX 파일 저장 실패 - 권한 오류")
            print("   해결 방법:")
            print("   1. 기존 DOCX 파일이 Word, 한글 등에서 열려있다면 닫아주세요")
            print("   2. output 폴더의 DOCX 파일을 삭제하거나 이름을 변경하세요")
            print("   3. 다시 시도해주세요")
            return None

        except Exception as e:
            print(f"\n⚠️ DOCX 생성 실패: {e}")
            import traceback
            traceback.print_exc()
            print("   TXT 파일로 대체 저장합니다.")
            return save_report(content, config, templates, 'txt', isopleth_images)

    elif format == 'hwp':
        print("\n📝 HWP 형식으로 저장하기 위해 먼저 DOCX를 생성합니다.")
        result = save_report(content, config, templates, 'docx', isopleth_images)
        print("\n   📌 HWP 파일이 필요하면:")
        print("      1. 생성된 DOCX 파일을 한글(HWP)에서 열기")
        print("      2. '파일 > 다른 이름으로 저장 > HWP 형식' 선택")
        return result
