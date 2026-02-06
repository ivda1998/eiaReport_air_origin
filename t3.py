import os
import time
import subprocess
import pandas as pd
from anthropic import Anthropic
try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None

# ==========================================
# 환경 설정
# ==========================================
# Claude API 키 설정 (환경변수 또는 직접 입력)
ANTHROPIC_API_KEY = os.environ.get("ANTHROPIC_API_KEY", "YOUR_API_KEY_HERE")  # 여기에 API 키 입력
if not ANTHROPIC_API_KEY:
    print("⚠️ ANTHROPIC_API_KEY 환경변수를 설정하거나 코드에서 직접 입력하세요.")
    print("   예: export ANTHROPIC_API_KEY='your-api-key-here'")

client = Anthropic(api_key=ANTHROPIC_API_KEY) if ANTHROPIC_API_KEY else None

BASE_DIR = r"G:\eia"
INPUT_DIR = os.path.join(BASE_DIR, "input")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
SAMPLE_DIR = os.path.join(BASE_DIR, "sample")
DATA_DIR = os.path.join(BASE_DIR, "data")
ENGINE_DIR = os.path.join(BASE_DIR, "engine")

AERMOD_EXE = os.path.join(ENGINE_DIR, "aermod.exe")
INPUT_FILE = os.path.join(INPUT_DIR, "project.inp")
OUTPUT_FILE = os.path.join(INPUT_DIR, "project.out")


# ==========================================
# 1. 참조 자료 로딩 함수
# ==========================================
def load_sample_reports():
    """sample 폴더의 PDF 보고서를 텍스트로 추출"""
    samples = []

    if not os.path.exists(SAMPLE_DIR):
        print(f"⚠️ 샘플 폴더를 찾을 수 없습니다: {SAMPLE_DIR}")
        return samples

    print(f"\n📂 샘플 보고서 로딩 중...")
    for filename in os.listdir(SAMPLE_DIR):
        if filename.endswith('.pdf'):
            filepath = os.path.join(SAMPLE_DIR, filename)
            try:
                if PdfReader is None:
                    print(f"  ⚠️ PyPDF2가 설치되지 않아 {filename}을 읽을 수 없습니다.")
                    continue

                reader = PdfReader(filepath)
                text = ""
                # 최대 50페이지까지만 읽기 (토큰 제한)
                max_pages = min(50, len(reader.pages))
                for page_num in range(max_pages):
                    text += reader.pages[page_num].extract_text()

                samples.append({
                    'filename': filename,
                    'content': text[:30000]  # 최대 30,000자
                })
                print(f"  ✓ {filename} 로딩 완료 ({len(reader.pages)} 페이지)")
            except Exception as e:
                print(f"  ✗ {filename} 로딩 실패: {e}")

    return samples


def load_facility_data():
    """data 폴더의 정온시설 위치 정보 로딩"""
    facility_file = os.path.join(DATA_DIR, "정온시설.xlsx")

    if not os.path.exists(facility_file):
        print(f"⚠️ 정온시설 파일을 찾을 수 없습니다: {facility_file}")
        return None

    try:
        df = pd.read_excel(facility_file)
        # 헤더 행 제외 (보통 처음 2행)
        if len(df) > 2:
            df = df.iloc[2:]  # 3번째 행부터 데이터

        # 실제 데이터 개수 확인
        actual_count = len(df)
        print(f"✓ 정온시설 데이터 로딩 완료 ({actual_count}개 시설)")

        # 전체 데이터를 문자열로 변환
        pd.set_option('display.max_colwidth', None)
        pd.set_option('display.max_rows', None)
        result = df.to_string(index=False, max_colwidth=100)
        pd.reset_option('display.max_colwidth')
        pd.reset_option('display.max_rows')

        return result
    except Exception as e:
        print(f"✗ 정온시설 데이터 로딩 실패: {e}")
        return None


def load_measurement_data():
    """data 폴더의 현황측정자료 로딩"""
    measurement_file = os.path.join(DATA_DIR, "현황측정자료.xlsx")

    if not os.path.exists(measurement_file):
        print(f"⚠️ 현황측정자료를 찾을 수 없습니다: {measurement_file}")
        return None

    try:
        df = pd.read_excel(measurement_file)
        print(f"✓ 현황측정자료 로딩 완료 ({len(df)}개 데이터)")
        return df.to_string(index=False)
    except Exception as e:
        print(f"✗ 현황측정자료 로딩 실패: {e}")
        return None


def load_reference_papers():
    """외부 전문자료 폴더에서 논문/보고서 로딩"""
    references_dir = os.path.join(DATA_DIR, "references")
    references = []

    if not os.path.exists(references_dir):
        print(f"ℹ️ references 폴더가 없습니다. 외부 전문자료 없이 진행합니다.")
        return None

    print(f"\n📚 외부 전문자료 로딩 중...")
    for filename in os.listdir(references_dir):
        if filename.endswith(('.txt', '.pdf', '.docx')):
            filepath = os.path.join(references_dir, filename)
            try:
                if filename.endswith('.txt'):
                    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                        content = f.read()
                        references.append({
                            'filename': filename,
                            'content': content[:5000]  # 최대 5,000자
                        })
                        print(f"  ✓ {filename} 로딩 완료")
                elif filename.endswith('.pdf') and PdfReader:
                    reader = PdfReader(filepath)
                    text = ""
                    max_pages = min(10, len(reader.pages))
                    for page_num in range(max_pages):
                        text += reader.pages[page_num].extract_text()
                    references.append({
                        'filename': filename,
                        'content': text[:5000]
                    })
                    print(f"  ✓ {filename} 로딩 완료")
            except Exception as e:
                print(f"  ✗ {filename} 로딩 실패: {e}")

    if references:
        print(f"✓ 총 {len(references)}개 외부 전문자료 로딩 완료")
        return "\n\n".join([
            f"[참고문헌: {item['filename']}]\n{item['content']}"
            for item in references
        ])
    return None


# ==========================================
# 2. AERMOD 실행 함수
# ==========================================
def run_aermod():
    """AERMOD 모델링 실행"""
    if not os.path.exists(AERMOD_EXE):
        return f"❌ AERMOD 실행 파일을 찾을 수 없습니다: {AERMOD_EXE}"

    if not os.path.exists(INPUT_FILE):
        return f"❌ 입력 파일을 찾을 수 없습니다: {INPUT_FILE}"

    print(f"\n🚀 AERMOD 시뮬레이션 시작...")
    print(f"   입력 파일: {INPUT_FILE}")

    try:
        # 작업 디렉토리를 input 폴더로 변경
        original_dir = os.getcwd()
        os.chdir(INPUT_DIR)

        # AERMOD 실행
        result = subprocess.run(
            [AERMOD_EXE, "project.inp"],
            capture_output=True,
            text=True,
            timeout=300  # 5분 타임아웃
        )

        os.chdir(original_dir)

        if result.returncode == 0:
            print("✅ AERMOD 모델링 완료")
            return "SUCCESS"
        else:
            print(f"⚠️ AERMOD 실행 중 경고 발생 (코드: {result.returncode})")
            return f"WARNING: {result.stderr[:500]}"

    except subprocess.TimeoutExpired:
        os.chdir(original_dir)
        return "❌ AERMOD 실행 시간 초과 (5분)"
    except Exception as e:
        os.chdir(original_dir)
        return f"❌ AERMOD 실행 오류: {str(e)}"


# ==========================================
# 3. 결과 분석 및 보고서 작성 (Claude API 사용)
# ==========================================
def analyze_output():
    """AERMOD 결과를 분석하여 환경영향평가서 작성 (Claude 사용)"""

    if not client:
        return "❌ Claude API 키가 설정되지 않았습니다. ANTHROPIC_API_KEY를 설정하세요."

    if not os.path.exists(OUTPUT_FILE):
        return "❌ AERMOD 출력 파일을 찾을 수 없습니다. 먼저 모델링을 실행하세요."

    print(f"\n📊 AERMOD 결과 분석 중...")

    # AERMOD 결과 파일 읽기
    with open(OUTPUT_FILE, 'r', encoding='utf-8', errors='ignore') as f:
        aermod_data = f.read()[-20000:]  # 마지막 20,000자

    # 참조 자료 로딩
    print(f"\n📚 참조 자료 로딩 중...")
    sample_reports = load_sample_reports()
    facility_data = load_facility_data()
    measurement_data = load_measurement_data()
    reference_papers = load_reference_papers()

    # 샘플 보고서 텍스트 결합
    sample_text = "\n\n".join([
        f"[샘플 보고서: {item['filename']}]\n{item['content']}"
        for item in sample_reports
    ]) if sample_reports else "샘플 보고서 없음"

    # Claude API 프롬프트 구성
    prompt = f"""당신은 환경영향평가 대기질 분야 전문가입니다.
제공된 AERMOD 모델링 결과와 참조자료를 바탕으로 환경영향평가서의 대기질 부분을 작성하세요.

[작성 지침]
1. 샘플 보고서의 양식과 구성을 참조하여 동일한 형식으로 작성
2. 정온시설 위치 정보를 활용하여 영향 예측 대상을 명시 (전체 38개 시설 모두 포함)
3. 현황측정자료를 참조하여 현황과 예측결과를 비교·분석
4. 오염물질별 최대 착지농도와 발생 위치(좌표)를 명확히 제시
5. 환경기준 대비 평가 결과 및 적부 판정
6. 문체는 "~함", "~임"을 사용하고 전문용어를 정확히 사용
7. 보고서로 바로 활용 가능하도록 체계적으로 작성
8. **중요**: 예측 결과 해석 시 다음 사항 반드시 포함:
   - 관련 선행연구나 논문의 연구결과를 인용하여 예측결과의 타당성 검증
   - 예: "이는 Kim et al.(2020)의 연구에서 제시한 도로변 대기질 영향범위와 유사한 결과임"
   - 예: "환경부(2021) 대기질 관리지침에 따르면..."
   - 예: "선행 유사사업(OO 도로건설, 2019)의 예측결과와 비교 시..."
   - 국내외 학술논문, 정부 보고서, 유사 환경영향평가서 등을 적극 인용
   - 인용 형식: 저자(연도) 또는 기관명(연도) 형식 사용

[참조 1: 샘플 보고서 양식]
{sample_text[:15000]}

[참조 2: 정온시설 위치 정보 (총 38개 시설)]
{facility_data[:10000] if facility_data else "정온시설 정보 없음"}

[참조 3: 현황측정자료]
{measurement_data[:5000] if measurement_data else "현황측정자료 없음"}

[참조 4: 외부 전문자료 및 선행연구]
{reference_papers[:5000] if reference_papers else "외부 전문자료가 없습니다. 대신 일반적인 대기질 관련 학술 연구와 정부 보고서를 인용하여 작성하세요. 예: 환경부(2022), 국립환경과학원(2021), Lee et al.(2020) 등"}

[AERMOD 모델링 결과]
{aermod_data}

위 정보를 종합하여 환경영향평가서 대기질 부분을 작성하세요.
**반드시 선행연구와 전문자료를 인용하여 예측결과의 신뢰성을 높이세요.**"""

    # Claude API 호출 (재시도 로직 포함)
    print("\n🤖 Claude AI 보고서 작성 중... (최대 2-3분 소요)")
    max_retries = 3

    for attempt in range(max_retries):
        try:
            response = client.messages.create(
                model="claude-sonnet-4-20250514",  # 최신 Claude 4 Sonnet 모델
                max_tokens=16000,  # 긴 보고서 생성
                temperature=0.3,   # 일관성 있는 출력
                messages=[
                    {
                        "role": "user",
                        "content": prompt
                    }
                ]
            )

            # 응답 검증
            if response and response.content:
                result_text = response.content[0].text.strip()
                print(f"✓ Claude 응답 생성 완료 ({len(result_text):,}자)")
                return result_text
            else:
                print("⚠️ Claude 응답이 비어있습니다.")
                if attempt < max_retries - 1:
                    print(f"   재시도 중... ({attempt+1}/{max_retries})")
                    time.sleep(10)
                    continue
                else:
                    raise Exception("Claude 응답이 생성되지 않았습니다.")

        except Exception as e:
            error_msg = str(e)
            if "rate_limit" in error_msg.lower() and attempt < max_retries - 1:
                wait_time = 60
                print(f"⏳ API 요청 제한 초과. {wait_time}초 대기 중... ({attempt+1}/{max_retries})")
                time.sleep(wait_time)
            elif attempt < max_retries - 1:
                print(f"⚠️ 오류 발생: {error_msg}")
                print(f"   재시도 중... ({attempt+1}/{max_retries})")
                time.sleep(10)
            else:
                print(f"❌ Claude 보고서 생성 실패: {error_msg}")
                raise Exception(f"Claude 보고서 생성 실패: {error_msg}")

    return None


def save_report(content, format='txt'):
    """보고서를 파일로 저장"""
    if not os.path.exists(OUTPUT_DIR):
        os.makedirs(OUTPUT_DIR)

    # 내용 검증
    if not content or not content.strip():
        print("⚠️ 저장할 내용이 비어있습니다.")
        return None

    if format == 'txt':
        report_file = os.path.join(OUTPUT_DIR, "대기질_환경영향평가서_Claude.txt")
        try:
            with open(report_file, "w", encoding="utf-8") as f:
                f.write(content)
            print(f"\n💾 보고서 저장 완료: {report_file}")
            print(f"   파일 크기: {len(content):,}자")
            return report_file
        except Exception as e:
            print(f"❌ TXT 파일 저장 실패: {e}")
            return None

    elif format == 'docx':
        if Document is None:
            print("\n⚠️ python-docx가 설치되지 않았습니다. TXT 파일로 저장합니다.")
            return save_report(content, 'txt')

        # 타임스탬프를 추가하여 파일명 충돌 방지
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        report_file = os.path.join(OUTPUT_DIR, f"대기질_환경영향평가서_Claude_{timestamp}.docx")

        try:
            doc = Document()

            # 문서 제목
            title = doc.add_heading('환경영향평가서 - 대기질 부문 (Claude AI 작성)', level=1)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 내용을 줄 단위로 처리
            lines = content.split('\n')
            i = 0

            while i < len(lines):
                line = lines[i].strip()

                if not line:
                    i += 1
                    continue

                # 표 감지 (| 로 시작하는 라인)
                if line.startswith('|'):
                    table_lines = []
                    while i < len(lines) and lines[i].strip().startswith('|'):
                        table_lines.append(lines[i].strip())
                        i += 1

                    if len(table_lines) > 1:
                        try:
                            header_cells = [cell.strip() for cell in table_lines[0].split('|')[1:-1]]
                            num_cols = len(header_cells)

                            data_lines = [table_lines[0]]
                            for tl in table_lines[1:]:
                                if not tl.replace('|', '').replace('-', '').replace(' ', '').replace(':', ''):
                                    continue
                                data_lines.append(tl)

                            if len(data_lines) > 0:
                                table = doc.add_table(rows=len(data_lines), cols=num_cols)
                                table.style = 'Light Grid Accent 1'

                                for row_idx, data_line in enumerate(data_lines):
                                    cells = [cell.strip() for cell in data_line.split('|')[1:-1]]
                                    for col_idx, cell_text in enumerate(cells):
                                        if col_idx < num_cols:
                                            cell = table.rows[row_idx].cells[col_idx]
                                            cell.text = cell_text
                                            if row_idx == 0:
                                                for paragraph in cell.paragraphs:
                                                    for run in paragraph.runs:
                                                        run.bold = True
                        except Exception as e:
                            p = doc.add_paragraph('\n'.join(table_lines))
                            p.paragraph_format.line_spacing = 1.0

                    continue

                # 마크다운 제목 처리
                if line.startswith('#'):
                    level = min(line.count('#', 0, 4), 3)
                    text = line.lstrip('#').strip()
                    if text:
                        doc.add_heading(text, level=level)
                    i += 1
                    continue

                # 일반 문단
                para_lines = [line]
                i += 1
                while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith('#') and not lines[i].strip().startswith('|'):
                    para_lines.append(lines[i].strip())
                    i += 1

                para_text = ' '.join(para_lines)
                if para_text:
                    p = doc.add_paragraph(para_text)
                    p.paragraph_format.line_spacing = 1.5

                    for run in p.runs:
                        run.font.name = '맑은 고딕'
                        run.font.size = Pt(10)

            doc.save(report_file)
            print(f"\n💾 보고서 저장 완료: {report_file}")
            print(f"   파일명: {os.path.basename(report_file)}")
            print("   ℹ️ DOCX 파일은 한글(HWP)에서 열어서 HWP 형식으로 저장할 수 있습니다.")
            return report_file

        except Exception as e:
            print(f"\n⚠️ DOCX 생성 실패: {e}")
            import traceback
            traceback.print_exc()
            print("   TXT 파일로 대체 저장합니다.")
            return save_report(content, 'txt')


# ==========================================
# 4. 메인 실행
# ==========================================
if __name__ == "__main__":
    print("\n" + "="*70)
    print("   환경영향평가 대기질 보고서 자동 작성 시스템 (Claude AI)")
    print("="*70)

    if not client:
        print("\n❌ Claude API를 사용하려면 ANTHROPIC_API_KEY를 설정해야 합니다.")
        print("   1. Anthropic 웹사이트에서 API 키 발급: https://console.anthropic.com/")
        print("   2. 환경변수 설정: export ANTHROPIC_API_KEY='your-api-key-here'")
        print("   3. 또는 코드 21번째 줄에 직접 입력")
        exit(1)

    while True:
        print("\n[명령어]")
        print("  1. 모델링 - AERMOD 시뮬레이션 실행")
        print("  2. 분석   - 결과 분석 및 보고서 작성 (TXT)")
        print("  3. 전체   - 모델링 + 분석 한번에 실행 (TXT)")
        print("  4. DOCX   - 보고서를 DOCX 형식으로 작성")
        print("  5. 종료   - 프로그램 종료")

        command = input("\n명령을 입력하세요: ").strip()

        if command == "종료" or command == "5":
            print("\n👋 프로그램을 종료합니다.")
            break

        elif command == "모델링" or command == "1":
            result = run_aermod()
            print(result)

        elif command == "분석" or command == "2":
            try:
                report = analyze_output()
                if report:
                    print("\n" + "="*70)
                    print("[Claude AI 작성 환경영향평가서 - 대기질]")
                    print("="*70)
                    if len(report) > 1000:
                        print(report[:1000])
                        print(f"\n... (총 {len(report):,}자, 이하 생략) ...\n")
                    else:
                        print(report)

                    result = save_report(report, format='txt')
                    if result:
                        print(f"✓ 보고서 작성 완료")
                else:
                    print("❌ 보고서 내용이 생성되지 않았습니다.")
            except Exception as e:
                print(f"\n❌ 오류 발생: {str(e)}")
                import traceback
                traceback.print_exc()

        elif command == "전체" or command == "3":
            result = run_aermod()
            print(result)

            if "SUCCESS" in result or "WARNING" in result:
                try:
                    report = analyze_output()
                    if report:
                        print("\n" + "="*70)
                        print("[Claude AI 작성 환경영향평가서 - 대기질]")
                        print("="*70)
                        if len(report) > 1000:
                            print(report[:1000])
                            print(f"\n... (총 {len(report):,}자, 이하 생략) ...\n")
                        else:
                            print(report)

                        result = save_report(report, format='txt')
                        if result:
                            print(f"✓ 보고서 작성 완료")
                    else:
                        print("❌ 보고서 내용이 생성되지 않았습니다.")
                except Exception as e:
                    print(f"\n❌ 보고서 작성 오류: {str(e)}")
                    import traceback
                    traceback.print_exc()
            else:
                print("\n❌ 모델링 실패로 분석을 진행할 수 없습니다.")

        elif command == "DOCX" or command == "docx" or command == "4":
            try:
                if os.path.exists(OUTPUT_FILE):
                    report = analyze_output()
                    if report:
                        print("\n" + "="*70)
                        print("[Claude AI 작성 환경영향평가서 - 대기질]")
                        print("="*70)
                        print(f"보고서 길이: {len(report):,}자")
                        print(report[:500] + "..." if len(report) > 500 else report)

                        result = save_report(report, format='docx')
                        if result:
                            print(f"✓ DOCX 파일 생성 완료")
                    else:
                        print("❌ 보고서 내용이 생성되지 않았습니다.")
                else:
                    print("\n❌ AERMOD 출력 파일이 없습니다. 먼저 '모델링' 또는 '전체'를 실행하세요.")
            except Exception as e:
                print(f"\n❌ DOCX 생성 오류: {str(e)}")
                import traceback
                traceback.print_exc()

        else:
            print("⚠️ 올바른 명령을 입력하세요.")
